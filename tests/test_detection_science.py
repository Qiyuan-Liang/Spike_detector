"""Scientific invariants, using generated signals rather than user recordings."""
import os
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
os.environ.setdefault('MPLCONFIGDIR', '/tmp/spike-detector-test-mpl')
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PyQt6 import QtWidgets
from spike_detector.gui import (MainWindow, StatsViewerDialog, DetectionViewerDialog,
    BatchDetectionDialog, DetectionSettingsDialog, SettingsDialog, DenoisingSettingsDialog)
from spike_detector.gui import estimate_noise_mad as gui_noise_mad
from spike_detector.utils.session import normalize_time_and_fs, load_session_path
from spike_detector.utils.batch import discover_batch_sessions
from spike_detector.utils.validation import validate_time, effective_settings, mask_windows
from spike_detector.utils.widths import measure_width, measure_widths
from spike_detector.utils.stats import autocorrelogram_counts
from spike_detector.utils.processing import (process_cell_simple, process_cell_template_matching,
    _build_parallel_template_banks, apply_filter, exclusion_mask,
    estimate_noise_mad as processing_noise_mad)


class ScienceTests(unittest.TestCase):
    def test_acg_counts_distinct_pairs_symmetrically_without_cross_recording_pairs(self):
        edges, counts = autocorrelogram_counts([0, 10, 30], 30, 10)
        np.testing.assert_array_equal(edges, [-30, -20, -10, 0, 10, 20, 30])
        np.testing.assert_array_equal(counts, [1, 1, 1, 1, 1, 1])
        _, second = autocorrelogram_counts([100, 120], 30, 10)
        np.testing.assert_array_equal(counts + second, [1, 2, 1, 1, 2, 1])
        _, duplicate = autocorrelogram_counts([0, 0, 10], 10, 10)
        np.testing.assert_array_equal(duplicate, [2, 2])

    def test_explicit_seconds_100hz_and_origin(self):
        _, ms, fs, valid = normalize_time_and_fs(7 + np.arange(100)*.01, time_unit='s')
        self.assertAlmostEqual(fs, 100)
        self.assertAlmostEqual(ms[0], 7000)
        self.assertTrue(valid.all())
        _, other, other_fs, _ = normalize_time_and_fs(ms, time_unit='ms')
        np.testing.assert_allclose(ms, other)
        self.assertAlmostEqual(fs, other_fs)

    def test_template_resampling_preserves_sample_time(self):
        from spike_detector.utils.processing import _resample_trace_to_fs
        original = np.arange(20)/1000
        resampled = _resample_trace_to_fs(original, 1000, 5000)
        np.testing.assert_allclose(resampled, np.arange(96)/5000)

    def test_bad_timestamps_rejected(self):
        for t in ([0, 1, 1], [2, 1, 0], [0, np.nan, 2], [0, np.inf, 2], [0, 1, 3]):
            with self.subTest(t=t), self.assertRaises(ValueError):
                validate_time(t)
        with self.assertRaisesRegex(ValueError, 'disagrees'):
            validate_time([0, 1, 2], fs=100)
        _, inferred, fs, _ = normalize_time_and_fs([0, 1], time_unit='auto')
        self.assertEqual(fs, 1000)
        np.testing.assert_array_equal(inferred, [0, 1])

    def test_auto_time_units_from_header_and_spacing(self):
        with tempfile.TemporaryDirectory() as folder:
            for header, start, step, expected_unit in (
                ('t(ms)', 0, .3, 'ms'),
                ('time_s', 7, .001, 's'),
                ('time', 7, .001, 's'),
            ):
                times = start + np.arange(160) * step
                path = Path(folder) / f'{expected_unit}_{header.replace("/", "_")}.csv'
                pd.DataFrame({header: times, 'cell': np.sin(np.arange(160))}).to_csv(path, index=False)
                data = load_session_path(str(path))
                self.assertEqual(data['input_time_unit'], expected_unit)
                self.assertAlmostEqual(data['time_ms'][0], times[0] * (1000 if expected_unit == 's' else 1))

    def test_auxiliary_and_short_recordings_are_rejected_before_detection(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            offsets = root / 'Tri_puff_10_time_offsets.csv'
            pd.DataFrame({'time_ms': np.arange(100), 'offset': np.arange(100)}).to_csv(offsets, index=False)
            with self.assertRaisesRegex(ValueError, 'Auxiliary time-offset'):
                load_session_path(str(offsets))
            coordinates = root / 'Tri_puff_10_coordinates.csv'
            coordinates.write_text('x,y\n1,2\n')
            with self.assertRaisesRegex(ValueError, 'Auxiliary time-offset/coordinate'):
                load_session_path(str(coordinates))
            short = root / 'short.csv'
            pd.DataFrame({'time_ms': np.arange(12), 'cell': np.arange(12)}).to_csv(short, index=False)
            with self.assertRaisesRegex(ValueError, 'at least 128 time samples'):
                load_session_path(str(short))
            short_npz = root / 'short.npz'
            np.savez(short_npz, time_ms=np.arange(12), raw_data=np.ones((12, 1)),
                     cell_names=['cell'], fs=1000.)
            with self.assertRaisesRegex(ValueError, 'at least 128 time samples'):
                load_session_path(str(short_npz))

    def test_similarity_detects_close_burst_with_short_core(self):
        fs = 3333.333333333333
        t = np.arange(1200) / fs
        waveform = np.exp(-.5 * ((np.arange(41) - 20) / 2.8) ** 2)
        x = np.random.default_rng(7).normal(0, .12, t.size)
        x[495:536] += 3 * waveform
        x[505:546] += 2 * waveform
        res = process_cell_template_matching(x, fs, negative_going=False,
            template_ss_bank=[waveform], template_ss_fs_bank=[fs],
            template_match_method='Normalized Similarity', ss_similarity_threshold=.7,
            ss_min_dist_ms=2, cs_high_cut=0, ss_high_cut=0,
            use_preprocessed=True, pre_detrended=x, pre_baseline=np.zeros_like(x))
        self.assertTrue(np.any(np.abs(res['ss_peaks'] - 515) <= 2))
        self.assertTrue(np.any(np.abs(res['ss_peaks'] - 525) <= 2))

    def test_parallel_groups_are_all_available_and_selection_changes_detector_score(self):
        fs = 5000.0
        t = np.arange(151) / fs
        narrow = np.exp(-.5*((t-.015)/.001)**2)
        broad = np.exp(-.5*((t-.015)/.003)**2)
        bank = [narrow, narrow*2, broad, broad*2]
        groups = _build_parallel_template_banks(bank, [fs]*4, fs, force_peak_positive=True,
                                                 max_use_types=2, n_components=1)
        self.assertEqual(len(groups), 2)
        second = _build_parallel_template_banks(bank, [fs]*4, fs, force_peak_positive=True,
                                                 max_use_types=2, n_components=1, selected_groups=[2])
        self.assertEqual(len(second), 1)
        np.testing.assert_allclose(np.mean(second[0][0], axis=0), np.mean(groups[1][0], axis=0))
        from spike_detector.utils.processing import _resample_template_to_fs
        groups_10k = _build_parallel_template_banks(bank, [fs]*4, 10000.0, force_peak_positive=True,
                                                     max_use_types=2, n_components=1)
        self.assertEqual([len(group[0]) for group in groups_10k], [len(group[0]) for group in groups])
        for group_5k, group_10k in zip(groups, groups_10k):
            np.testing.assert_allclose(group_5k[0][0],
                _resample_template_to_fs(group_10k[0][0], 10000.0, fs), atol=1e-3)
        trace = np.random.default_rng(7).normal(0, .02, 3000)
        trace[1350:1501] += 2*broad
        kwargs = dict(template_cs_bank=bank, template_cs_fs_bank=[fs]*4,
                      negative_going=False, cs_high_cut=0, ss_high_cut=0,
                      cs_min_fwhm_ms=0, initial_blank_ms=0,
                      parallel_match=True, parallel_groups=2, parallel_components=1,
                      use_preprocessed=True, pre_detrended=trace, pre_baseline=np.zeros_like(trace))
        first_result = process_cell_template_matching(trace, fs, cs_selected_groups=[1], **kwargs)
        second_result = process_cell_template_matching(trace, fs, cs_selected_groups=[2], **kwargs)
        self.assertFalse(np.allclose(first_result['cs_similarity_trace'], second_result['cs_similarity_trace']))

    def test_template_lowpass_rejects_narrowband_artifact(self):
        fs = 3333.333333333333
        t = np.arange(3000) / fs
        artifact = np.sin(2 * np.pi * 1111 * t)
        res = process_cell_template_matching(artifact, fs, negative_going=False,
            cs_high_cut=0, ss_high_cut=0, template_ss_lowpass_hz=700,
            use_preprocessed=True, pre_detrended=artifact, pre_baseline=np.zeros_like(artifact))
        self.assertLess(np.std(res['ss_trace'][100:-100]), .05)
        self.assertEqual(res['template_ss_lowpass_hz_used'], 700)

    def test_larger_neighbor_cannot_set_width(self):
        t = np.arange(1501)/10
        trace = 3*np.exp(-.5*((t-60)/.8493)**2) + 8*np.exp(-.5*((t-70)/2.1233)**2)
        rows = measure_widths(trace, [600, 700], 10000, 50)
        self.assertAlmostEqual(rows[0]['fwhm_ms'], 2, delta=.05)
        self.assertAlmostEqual(rows[1]['fwhm_ms'], 5, delta=.2)
        self.assertEqual(rows[0]['peak_index'], 600)

    def test_candidate_not_global_max_even_without_neighbor_list(self):
        t = np.arange(1001)/10
        x = np.exp(-.5*((t-50)/.85)**2) + 5*np.exp(-.5*((t-60)/2.12)**2)
        row = measure_width(x, 500, 10000)
        self.assertAlmostEqual(row['fwhm_ms'], 2, delta=.05)

    def test_overlap_and_clipping_flagged(self):
        x = np.zeros(200)
        x[90:111] = 4
        x[100] = 5
        x[108] = 6
        row = measure_width(x, 100, 1000, neighbors=[100,108])
        self.assertTrue(np.isnan(row['fwhm_ms']))
        self.assertEqual(row['reason'], 'overlapping_events')
        x = np.zeros(100)
        x[:8] = [3,4,5,4,3,2,1,0]
        row = measure_width(x, 2, 1000)
        self.assertEqual(row['status'], 'uncertain')

    def test_template_alignment_is_bounded(self):
        x = np.zeros(100)
        x[48:53] = [1,3,5,3,1]
        x[60] = 100
        row = measure_width(x, 49, 1000, align_ms=2)
        self.assertEqual(row['peak_index'], 50)
        self.assertTrue(np.isfinite(row['fwhm_ms']))
        row = measure_width(x, 40, 1000, align_ms=2)
        self.assertEqual(row['status'], 'uncertain')
        low_rate = np.zeros(120)
        low_rate[49:52] = [1, 3, 1]
        row = measure_width(low_rate, 49, 450, align_ms=2)
        self.assertEqual(row['peak_index'], 50)
        self.assertTrue(np.isfinite(row['fwhm_ms']))

    def test_masks_have_explicit_sides(self):
        mask = exclusion_mask(100, 1000, [50], pre_ms=2, post_ms=5, initial_ms=3)
        np.testing.assert_array_equal(np.flatnonzero(mask), np.r_[0:3,48:56])
        self.assertEqual(mask_windows({'SS_BLANK_MS':18}), (9,9))
        self.assertEqual(mask_windows({'SS_BLANK_MS':18,'SS_MASK_PRE_MS':0,'SS_MASK_POST_MS':18}), (0,18))
        self.assertFalse(exclusion_mask(100,1000,[50]).any())

    def test_mask_does_not_modify_trace_or_create_boundary_spikes(self):
        t = np.arange(600)
        x = 12*np.exp(-.5*((t-300)/4)**2) + np.exp(-.5*((t-280)/.8)**2)
        kwargs = dict(negative_going=False, cs_low_cut=0, cs_high_cut=0, ss_low_cut=0, ss_high_cut=0,
            use_preprocessed=True, pre_detrended=x, pre_baseline=np.zeros_like(x), cs_thresh_sigma=20,
            cs_min_fwhm_ms=4, ss_mask_pre_ms=5, ss_mask_post_ms=15)
        res = process_cell_simple(x,1000,**kwargs)
        np.testing.assert_array_equal(res['ss_trace'], x)
        self.assertIn(300, res['cs_peaks'])
        self.assertFalse(np.any(res['ss_exclusion_mask'][res['ss_peaks']]))
        self.assertNotIn(295, res['ss_peaks'])
        self.assertNotIn(316, res['ss_peaks'])

    def test_template_mask_and_waveform_width(self):
        t = np.arange(1000)
        x = np.exp(-.5*((t-500)/3)**2)
        tpl = np.exp(-.5*((np.arange(41)-20)/3)**2)
        res = process_cell_template_matching(x,1000,negative_going=False,
            cs_high_cut=0, template_cs_bank=[tpl],template_cs_fs_bank=[1000],
            template_ss_bank=[tpl],template_ss_fs_bank=[1000],
            cs_thresh_sigma=1, ss_thresh_sigma=1, cs_min_fwhm_ms=0,
            use_preprocessed=True,pre_detrended=x,pre_baseline=np.zeros_like(x))
        self.assertGreater(len(res['cs_peaks']),0)
        row = min(res['cs_width_candidates'],key=lambda r:abs(r['candidate_index']-500))
        self.assertAlmostEqual(row['fwhm_ms'],7.065,delta=.15)
        np.testing.assert_array_equal(res['ss_trace'],x)
        self.assertFalse(np.any(res['ss_exclusion_mask'][res['ss_peaks']]))

    def test_invalid_filter_and_effective_windows(self):
        x = np.ones(100)
        for low,high in [(10,5),(0,500),(-1,10),(1000,0)]:
            with self.subTest(low=low,high=high), self.assertRaises(ValueError):
                apply_filter(x,1000,low,high)
        eff = effective_settings({'CS_HIGH_CUT_HZ':0}, {'method':'Median','window_ms':10000},0,1000,100)
        self.assertEqual(eff['baseline_window_samples'],50)
        with self.assertRaises(ValueError):
            effective_settings({'CS_HIGH_CUT_HZ':150},{'method':'Median'},0,100,100)

    def test_local_noise_excludes_masked_values(self):
        from spike_detector.utils.processing import _masked_local_sigma
        rng=np.random.default_rng(17)
        x=rng.normal(size=500)
        mask=exclusion_mask(500,1000,[250],pre_ms=20,post_ms=30)
        first=_masked_local_sigma(x,mask,51)
        x[mask]=1e9
        np.testing.assert_allclose(first,_masked_local_sigma(x,mask,51))
        self.assertTrue(np.all(first>0))

    def test_nonfinite_trace_not_repaired_and_npz_fs_checked(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'bad.csv'
            path.write_text('time,cell\n0,1\n1,\n2,3\n')
            with self.assertRaisesRegex(ValueError,'nonfinite'):
                load_session_path(str(path))
            npz=Path(folder)/'bad.npz'
            np.savez(npz,time_ms=[0,1,2],raw_data=np.ones((3,1)),cell_names=['a'],fs=100)
            with self.assertRaisesRegex(ValueError,'disagrees'):
                load_session_path(str(npz))


class ScienceGuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app=QtWidgets.QApplication.instance() or QtWidgets.QApplication([])

    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.win=MainWindow()

    def tearDown(self):
        self.win.close()
        self.temp.cleanup()

    def test_setting_inputs_and_spin_edit_fields_have_hover_descriptions(self):
        data = {'raw_data': np.zeros((300, 1)), 'time_ms': np.arange(300),
                'cell_names': ['Cell'], 'fs': 1000.0, 'results': [None]}
        dialogs = [
            BatchDetectionDialog(self.win), DetectionSettingsDialog(self.win),
            DetectionViewerDialog(data, self.win), StatsViewerDialog(data, self.win),
            SettingsDialog(self.win), DenoisingSettingsDialog(None),
        ]
        try:
            for owner in [self.win, *dialogs]:
                for widget in owner.findChildren(QtWidgets.QWidget):
                    editable_text = isinstance(widget, QtWidgets.QPlainTextEdit) and not widget.isReadOnly()
                    input_widget = isinstance(widget, (QtWidgets.QAbstractSpinBox,
                        QtWidgets.QComboBox, QtWidgets.QCheckBox, QtWidgets.QSlider))
                    if input_widget or editable_text:
                        with self.subTest(window=type(owner).__name__, widget=type(widget).__name__):
                            self.assertTrue(widget.toolTip().strip())
                            if isinstance(widget, QtWidgets.QAbstractSpinBox):
                                self.assertEqual(widget.lineEdit().toolTip(), widget.toolTip())
        finally:
            for dialog in dialogs:
                dialog.close()

    def test_excel_untimed_tail_loads_timed_prefix_with_explicit_warning(self):
        path = Path(self.temp.name) / 'mixed_length.xlsx'
        timed = pd.DataFrame({
            't(ms)': np.arange(160) * 0.12,
            'cell_A': np.sin(np.arange(160)),
            'cell_B': np.cos(np.arange(160)),
        })
        timed.to_excel(path, sheet_name='Sheet1', index=False)
        book = load_workbook(path)
        sheet = book['Sheet1']
        for value in range(12):
            sheet.append([None, None, float(value)])
        book.save(path)

        data = load_session_path(str(path))
        self.assertEqual(data['raw_data'].shape, (160, 2))
        self.assertAlmostEqual(data['fs'], 1000 / 0.12)
        self.assertEqual(data['ignored_untimed_tail_rows'], 12)
        self.assertIn('cell_B: 12', data['load_warnings'][0])
        self.win.master_folder = self.temp.name
        self.win.refresh_sessions()
        self.assertIn('Input warning: Excluded 12 rows', self.win.text_stats.toPlainText())
        _, _, _, issues, _ = discover_batch_sessions([self.temp.name])
        self.assertTrue(any('Input warning:' in issue and 'cell_B: 12' in issue for issue in issues))

        sheet.cell(row=80, column=1).value = None
        book.save(path)
        with self.assertRaisesRegex(ValueError, 'finite samples'):
            load_session_path(str(path))

    def test_excel_lock_files_are_ignored_but_corrupt_workbooks_still_reported(self):
        root = Path(self.temp.name)
        real = root / 'BestSS.xlsx'
        pd.DataFrame({'time_ms': np.arange(160), 'cell': np.sin(np.arange(160))}).to_excel(real, index=False)
        locks = [root/'~$BestSS.xlsx', root/'~$BestSS_ASAP7.XLSX']
        for path in locks:
            path.write_bytes(b'Excel owner file, not a ZIP archive')
        bad = root/'zz_bad.xlsx'
        bad.write_bytes(b'corrupt real workbook')
        self.assertEqual(self.win._list_table_files_in_dir(str(root)), [str(real), str(bad)])
        self.assertEqual(load_session_path(str(root))['session_path'], str(real))
        for path in locks:
            with patch('spike_detector.utils.session.pd.read_excel') as reader:
                with self.assertRaisesRegex(ValueError, 'lock file'):
                    load_session_path(str(path))
                reader.assert_not_called()
        self.win.master_folder = str(root)
        self.win.refresh_sessions()
        self.assertEqual(self.win.sessions, [str(real)])
        self.assertIn('skipped invalid: 1', self.win.statusBar().currentMessage())
        self.assertNotIn('~$', self.win.text_stats.toPlainText())
        paths, _, _, issues, _ = discover_batch_sessions([str(root)])
        self.assertEqual(paths, [str(real.resolve())])
        self.assertEqual(len(issues), 1)
        self.assertIn('zz_bad.xlsx', issues[0])
        self.assertNotIn('~$', issues[0])
        self.assertTrue(all(path.exists() for path in locks))

    def test_save_load_restores_all_setting_groups_and_template_sources(self):
        root = Path(self.temp.name)
        cs_file = root / 'cs_templates.npz'
        ss_file = root / 'ss_templates.npz'
        np.savez(cs_file, template_cs=np.arange(21, dtype=float), fs_cs=1000.0)
        np.savez(ss_file, template_ss=np.arange(11, dtype=float), fs_ss=1000.0)
        for kind, source in (('CS', cs_file), ('SS', ss_file)):
            self.win._merge_templates_for_type(kind, self.win._read_template_npz(source), str(source))

        self.win.spin_cs_order.setValue(5)
        self.win.spin_ss_order.setValue(6)
        self.win.spin_cs_low.setValue(12)
        self.win.spin_cs_high.setValue(120)
        self.win.spin_cs_thresh.setValue(8)
        self.win.spin_ss_low.setValue(10)
        self.win.spin_ss_high.setValue(500)
        self.win.spin_ss_thresh.setValue(4.5)
        self.win.chk_template_parallel.setChecked(True)
        self.win.chk_template_whitening.setChecked(True)
        self.win.chk_combo.setChecked(True)
        self.win.spin_template_groups.setValue(4)
        self.win.combo_template_method.setCurrentIndex(1)
        self.win.spin_template_cs_sigma.setValue(9)
        self.win.qc_inputs['CS_MIN_AMPLITUDE_SNR'].setValue(4)
        self.win.qc_inputs['CS_MAX_FWHM_MS'].setValue(45)
        self.win.spin_template_ss_sigma.setValue(7)
        self.win.spin_template_cs_sim.setValue(0.85)
        self.win.spin_template_ss_sim.setValue(0.75)
        self.win.qc_inputs['SS_MIN_AMPLITUDE_SNR'].setValue(3)
        self.win.qc_inputs['SS_MIN_FWHM_MS'].setValue(.4)
        self.win.spin_template_lowpass.setValue(650)
        self.win.chk_local_baseline.setChecked(True)
        self.win.spin_local_ss_ms.setValue(60)
        self.win.spin_local_cs_ms.setValue(220)
        self.win.chk_detection_save.setChecked(True)
        self.win.combo_baseline_method.setCurrentText('Percentile')
        self.win.spin_baseline_window.setValue(35)
        self.win.spin_baseline_percentile.setValue(31)
        self.win.combo_frame_processing.setCurrentText('Mean')
        self.assertEqual(self.win._get_frame_processing_mode(), 'Downsampling frames')
        self.win.spin_avg_frames.setValue(3)
        self.win.slider_time.setValue(200)
        self.win.spin_window.setValue(750)
        self.win.spin_zoom.setValue(8)
        self.win.chk_show_baseline.setChecked(True)
        self.win.chk_show_cs.setChecked(True)
        self.win.tabs_detection.setCurrentIndex(1)
        self.win.chk_denoise.setChecked(True)
        advanced = SettingsDialog(self.win)
        self.win.tabs_polarity.setCurrentIndex(0)
        advanced.chk_denoise_cs.setChecked(True)
        advanced.spin_cs_mind.setValue(31)
        self.win.qc_inputs['CS_MIN_FWHM_MS'].setValue(7)
        advanced.spin_ss_mind.setValue(5)
        advanced.spin_ss_mask_pre.setValue(7)
        advanced.spin_ss_mask_post.setValue(13)
        self.win.qc_inputs['SS_MAX_FWHM_MS'].setValue(5.5)
        advanced.spin_initial_blank.setValue(180)
        advanced.spin_tpl_cs_window.setValue(40)
        advanced.spin_tpl_ss_window.setValue(12)
        advanced.spin_template_components.setValue(3)
        advanced.spin_line_scale.setValue(1.3)
        advanced.combo_raw_scale_unit.setCurrentIndex(1)
        advanced.save_and_close()
        advanced.close()
        denoise = DenoisingSettingsDialog(self.win)
        denoise._preview_timer.stop()
        denoise.spin_preview_window.setValue(300)
        denoise.spin_f_min.setValue(4)
        denoise.spin_f_max.setValue(900)
        denoise.spin_n_freqs.setValue(64)
        denoise.spin_min_clusters.setValue(3)
        denoise.spin_max_clusters.setValue(15)
        denoise.spin_pca.setValue(12)
        denoise.spin_thr_sigma.setValue(2.5)
        denoise.spin_att_min.setValue(0.4)
        denoise.chk_soft.setChecked(True)
        denoise.spin_cycles.setValue(1.5)
        denoise.spin_max_tp.setValue(15000)
        denoise.chk_event_refine.setChecked(False)
        denoise.spin_evt_win.setValue(25)
        denoise.spin_evt_sigma.setValue(3)
        denoise.spin_evt_z.setValue(-1.2)
        denoise.spin_evt_att.setValue(0.7)
        denoise.save_and_accept()
        denoise._preview_timer.stop()
        denoise.close()
        stats_data = {'raw_data': np.zeros((300, 1)), 'time_ms': np.arange(300),
                      'cell_names': ['Cell'], 'fs': 1000.0, 'results': [None]}
        stats = StatsViewerDialog(stats_data, self.win)
        stats.spin_cs_acg_window.setValue(800)
        stats.spin_cs_acg_bin.setValue(20)
        stats.spin_ss_acg_window.setValue(80)
        stats.spin_ss_acg_bin.setValue(2)
        stats.spin_cs_bg_alpha.setValue(0.4)
        stats.spin_ss_bg_alpha.setValue(0.004)
        stats.close()
        self.win.params.update({
            'TEMPLATE_CS_SELECTED_GROUPS': [1, 3],
        })
        self.win.colors['cs_trace'] = '#123456'
        expected = self.win._current_settings_snapshot()
        expected_qc = {key: widget.value() for key, widget in self.win.qc_inputs.items()}
        path = root / 'settings.json'
        with patch('spike_detector.gui.QFileDialog.getSaveFileName', return_value=(str(path), '')), \
             patch('spike_detector.gui.QMessageBox.information'):
            self.win.save_all_settings()
        self.assertEqual(json.loads(path.read_text()), expected)

        for widget in self.win.qc_inputs.values():
            widget.setValue(0)
        self.win.spin_cs_order.setValue(2)
        self.win.spin_ss_order.setValue(2)
        self.win.chk_template_parallel.setChecked(False)
        self.win.chk_template_whitening.setChecked(False)
        self.win.chk_combo.setChecked(False)
        self.win.spin_template_groups.setValue(2)
        self.win.slider_time.setValue(0)
        self.win.spin_baseline_percentile.setValue(20)
        self.win.chk_denoise.setChecked(False)
        self.win.params.update({
            'SS_MASK_PRE_MS': 2.0, 'SS_MASK_POST_MS': 2.0,
            'DENOISE_THRESHOLD_SIGMA': 1.0, 'DENOISE_SOFT_THRESHOLD': False,
            'STATS_CS_ACG_WINDOW_MS': 1000.0, 'STATS_SS_ACG_BIN_MS': 1.0,
            'RAW_SCALE_BAR_UNIT': 'Sigma', 'LINE_WIDTH_SCALE': 0.7,
        })
        self.win.colors['cs_trace'] = '#ffffff'
        self.win.clear_templates()
        with patch('spike_detector.gui.QFileDialog.getOpenFileName', return_value=(str(path), '')), \
             patch('spike_detector.gui.QMessageBox.information') as loaded, \
             patch('spike_detector.gui.QMessageBox.critical') as failed:
            self.win.load_all_settings()
        failed.assert_not_called()
        loaded.assert_called_once()
        self.assertEqual(self.win._current_settings_snapshot(), expected)
        self.assertEqual({key: widget.value() for key, widget in self.win.qc_inputs.items()}, expected_qc)
        self.assertEqual(len(self.win.template_store['cs_templates']), 1)
        self.assertEqual(len(self.win.template_store['ss_templates']), 1)
        advanced = SettingsDialog(self.win)
        self.assertEqual(advanced.spin_cs_mind.value(), 31)
        self.assertEqual(advanced.spin_ss_mask_post.value(), 13)
        self.assertEqual(advanced.combo_raw_scale_unit.currentData(), 'dF/F (%)')
        advanced.close()
        denoise = DenoisingSettingsDialog(self.win)
        denoise._preview_timer.stop()
        self.assertEqual(denoise.spin_thr_sigma.value(), 2.5)
        self.assertTrue(denoise.chk_soft.isChecked())
        denoise.close()
        stats = StatsViewerDialog(stats_data, self.win)
        self.assertEqual(stats.spin_ss_acg_bin.value(), 2)
        self.assertEqual(stats.spin_cs_bg_alpha.value(), 0.4)
        stats.close()
        cs_file.unlink()
        with patch('spike_detector.gui.QFileDialog.getOpenFileName', return_value=(str(path), '')), \
             patch('spike_detector.gui.QMessageBox.information') as loaded:
            self.win.load_all_settings()
        self.assertIn('Could not restore template sources', loaded.call_args.args[2])
        self.assertEqual(len(self.win.template_store['cs_templates']), 0)
        self.assertEqual(len(self.win.template_store['ss_templates']), 1)

    def test_detection_viewer_sigma_bar_uses_normal_scaled_mad(self):
        signal = np.array([0.0, 0.0, 0.0, 1.0])
        self.assertAlmostEqual(gui_noise_mad(signal), float(np.std(signal)))
        self.assertAlmostEqual(processing_noise_mad(signal), float(np.std(signal)))
        data = {'raw_data': np.zeros((300, 1)), 'time_ms': np.arange(300),
                'cell_names': ['Cell'], 'fs': 1000.0, 'results': [None]}
        viewer = DetectionViewerDialog(data, self.win)
        try:
            viewer.ax_raw_bot.clear()
            viewer._draw_classic_scale_bar(viewer.ax_raw_bot, 500, (0, 10), 2.0)
            label = viewer.ax_raw_bot.texts[-1].get_text()
            self.assertIn(r'\sigma_{\mathrm{MAD}}', label)
            bar_fraction = np.ptp(viewer.ax_raw_bot.lines[0].get_ydata())
            self.assertAlmostEqual(bar_fraction * 10, float(label.split()[0]) * 2.0)
        finally:
            viewer.close()

    def test_folder_skips_auxiliary_tables_and_failed_viewer_cell_is_safe(self):
        root = Path(self.temp.name)
        t = np.arange(256)
        pd.DataFrame({'time_ms': t, 'cell': 100 - np.sin(t / 5)}).to_csv(root / 'trace.csv', index=False)
        pd.DataFrame({'time_ms': t, 'offset': t * .1}).to_csv(
            root / 'Tri_puff_10_time_offsets.csv', index=False)
        pd.DataFrame({'time_ms': np.arange(12), 'cell': np.arange(12)}).to_csv(
            root / 'short.csv', index=False)
        self.win.master_folder = self.temp.name
        self.win.refresh_sessions()
        self.assertEqual(len(self.win.sessions), 1)
        self.assertIn('skipped invalid: 2', self.win.statusBar().currentMessage())
        batch_sessions, _, _, batch_issues, _ = discover_batch_sessions([str(root)])
        self.assertEqual(len(batch_sessions), 1)
        self.assertEqual(sum(issue.startswith('Skipped:') for issue in batch_issues), 2)
        self.assertEqual(self.win.run_detection_all(), (1, 1))
        self.win.data['results'][0] = None
        viewer = DetectionViewerDialog(self.win.data, self.win)
        viewer._on_cell_changed(0)
        self.assertIn('Detection unavailable', viewer.lbl_candidate_info.text())
        self.assertFalse(viewer.btn_polarity_control.isEnabled())
        viewer.close()

    def test_origin_width_flags_masks_and_exports(self):
        t=np.arange(1000)
        x=100-np.exp(-.5*((t-300)/.85)**2)*5
        pd.DataFrame({'time':10+t/1000,'cell':x}).to_csv(Path(self.temp.name)/'trace.csv',index=False)
        self.win.master_folder=self.temp.name
        self.win.refresh_sessions()
        self.win.spin_cs_high.setValue(0)
        self.win.params['CS_MIN_FWHM_MS']=100
        self.win.params['INITIAL_BLANK_MS']=0
        self.assertEqual(self.win.run_detection_all(),(1,1))
        data=self.win.data
        res=data['results'][0]
        self.assertIn(300,res['ss_peaks'])
        np.testing.assert_allclose(data['spike_times_ss'][0],data['time_ms'][res['ss_peaks']])
        self.assertGreater(data['spike_times_ss'][0][0],10000)
        with np.load(data['results_file'],allow_pickle=True) as saved:
            widths=np.asarray(saved['event_fwhm_ss'][0],dtype=float)
            self.assertEqual(len(widths),len(res['ss_peaks']))
            self.assertIn('width_quality_json',saved.files)
            self.assertEqual(saved['ss_exclusion_mask'].shape,(1000,1))
            settings=json.loads(str(saved['analysis_settings_json']))
            self.assertEqual(settings['input_time_unit'],'s')
        frozen=self.win._get_waveform_source_trace_for_stats(data,0).copy()
        self.win.spin_avg_frames.setValue(10)
        np.testing.assert_array_equal(self.win._get_waveform_source_trace_for_stats(data,0),frozen)
        dialog=StatsViewerDialog(data,self.win)
        self.assertTrue(any('Uncertain widths:' in txt.get_text() for ax in dialog.fig.axes for txt in ax.texts))
        self.assertEqual(len(dialog.fig.axes), 12)
        self.assertEqual(dialog.fig.axes[1].get_title(), 'CS ACG')
        self.assertEqual(dialog.fig.axes[7].get_title(), 'SS ACG')
        self.assertEqual(dialog.fig.axes[1].get_xlim(), (-1000.0, 1000.0))
        self.assertEqual(dialog.fig.axes[7].get_xlim(), (-100.0, 100.0))
        dialog.spin_ss_acg_window.setValue(80)
        dialog.spin_ss_acg_bin.setValue(2)
        self.assertEqual(dialog.fig.axes[7].get_xlim(), (-80.0, 80.0))
        self.assertEqual(dialog.spin_ss_acg_bin.value(), 2)
        self.assertEqual(self.win.params['STATS_SS_ACG_WINDOW_MS'], 80)
        self.assertEqual(self.win.params['STATS_SS_ACG_BIN_MS'], 2)
        dialog.close()

    def test_uncertain_not_silently_passed(self):
        trace=np.zeros(200)
        trace[100]=1
        trace[101]=2 # supplied candidate is not a local max
        data={'raw_data':trace[:,None],'time_ms':np.arange(200),'cell_names':['c'],'fs':1000}
        self.win.baseline_params['method']='Disable'
        self.win.params.update(INITIAL_BLANK_MS=0, FINAL_BLANK_MS=0)
        self.win.params['NEGATIVE_GOING']=False
        res={'ss_peaks':np.array([100]),'cs_peaks':np.array([],dtype=int),'cs_width_candidates':[]}
        self.win._apply_ss_fwhm_filter(res,data,0)
        np.testing.assert_array_equal(res['ss_peaks'],[100])
        self.assertEqual(res['ss_fwhm_filter_uncertain'],1)
        self.assertEqual(res['event_widths_ss'][0]['decision'],'uncertain')

    def test_viewer_marks_retained_cs_even_when_width_is_uncertain(self):
        t = np.arange(300)
        pd.DataFrame({'time_ms': t, 'cell': 100 + np.sin(t / 10)}).to_csv(
            Path(self.temp.name) / 'trace.csv', index=False)
        self.win.master_folder = self.temp.name
        self.win.refresh_sessions()
        self.assertEqual(self.win.run_detection_all(), (1, 1))
        result = self.win.data['results'][0]
        result['det_method'] = 'Template Matching (Amplitude-fit LLR)'
        self.win.params['DETECTION_VIEW_AUTO_START'] = False
        result['cs_peaks'] = np.array([50, 100])
        result['cs_candidate_diagnostics'] = dict(score_peaks=2, masked=0,
            response_rejected=0, peak_rejected=0, refractory_rejected=0)
        result['ss_candidate_diagnostics'] = dict(no_templates=True)
        result['event_widths_cs'] = [
            dict(candidate_index=50, status='measured', fwhm_ms=4.0,
                 left_index=48.0, right_index=52.0),
            dict(candidate_index=100, status='uncertain', fwhm_ms=np.nan,
                 reason='not_a_local_peak'),
        ]
        viewer = DetectionViewerDialog(self.win.data, self.win)
        markers = [line for line in viewer.ax_raw_bot.lines
                   if line.get_color() == self.win.colors.get('cs_trace', '#009E73')]
        self.assertEqual(len(markers), 2)
        self.assertEqual(sum(line.get_marker() == '_' for line in markers), 1)
        self.assertIn('width uncertain 1', viewer.lbl_candidate_info.text())
        viewer.close()

    def test_template_matching_retains_narrow_spike_without_minimum_width_gate(self):
        fs = 3333.333333333333
        t = np.arange(1000) / fs
        x = np.exp(-.5*((t-.10)/.00013)**2) + np.exp(-.5*((t-.20)/.00065)**2)
        data = {'raw_data': x[:,None], 'time_ms': t*1000, 'cell_names': ['c'], 'fs': fs}
        self.win.params['NEGATIVE_GOING'] = False
        self.win.baseline_params['method'] = 'Disable'
        self.win.spin_avg_frames.setValue(0)
        res = {'det_method': 'Template Matching (Normalized Similarity)',
               'ss_trace': x, 'ss_peaks': np.array([333,667]),
               'cs_peaks': np.array([],dtype=int), 'cs_width_candidates': []}
        self.win._apply_ss_fwhm_filter(res, data, 0)
        self.assertEqual(len(res['ss_peaks']), 2)
        self.assertEqual(res['ss_width_source'], 'filtered SS detection trace')
        self.assertEqual([row['decision'] for row in res['ss_width_candidates']], ['pass', 'pass'])

    def test_ss_filter_rejects_broad_neighbor_only(self):
        t=np.arange(1501)/10
        trace=3*np.exp(-.5*((t-60)/.8493)**2)+8*np.exp(-.5*((t-70)/2.1233)**2)
        data={'raw_data':trace[:,None],'time_ms':t,'cell_names':['c'],'fs':10000}
        self.win.baseline_params['method']='Disable'
        self.win.params.update(INITIAL_BLANK_MS=0, FINAL_BLANK_MS=0)
        self.win.params['NEGATIVE_GOING']=False
        res={'ss_peaks':np.array([600,700]),'cs_peaks':np.array([],dtype=int),'cs_width_candidates':[]}
        self.win._apply_ss_fwhm_filter(res,data,0)
        np.testing.assert_array_equal(res['ss_peaks'],[600])
        self.assertEqual([row['decision'] for row in res['ss_width_candidates']],['pass','fail'])

    def test_validation_failure_has_no_results_file(self):
        t=np.arange(1000)
        pd.DataFrame({'time':t,'cell':100+np.sin(t)}).to_csv(Path(self.temp.name)/'trace.csv',index=False)
        self.win.master_folder=self.temp.name
        self.win.refresh_sessions()
        self.win.spin_cs_high.setValue(600)
        self.assertEqual(self.win.run_detection_all(),(0,0))
        self.assertIn('Nyquist',self.win.text_stats.toPlainText())
        self.assertNotIn('results_file',self.win.data)
        self.assertEqual(self.win.data['results'],[None])


if __name__ == '__main__':
    unittest.main()
